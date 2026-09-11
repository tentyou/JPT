from pathlib import Path
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "competition"


def markdown_to_docx(source: Path, target: Path) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif len(line) > 2 and line[0].isdigit() and line[1:3] == ". ":
            doc.add_paragraph(line[3:], style="List Number")
        elif not line.startswith(("```", "|", "<http")):
            doc.add_paragraph(line.replace("**", "").replace("`", ""))
    doc.save(target)


def add_slide(prs: Presentation, title: str, bullets: list[str], accent=(15, 118, 110)) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = RGBColor(246, 248, 247)
    title_box = slide.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.8), Inches(.8))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title; title_p.font.name = "Microsoft YaHei"; title_p.font.size = PptPt(28)
    title_p.font.bold = True; title_p.font.color.rgb = RGBColor(*accent)
    body = slide.shapes.add_textbox(Inches(.85), Inches(1.55), Inches(11.5), Inches(5.5)).text_frame
    body.clear()
    for index, text in enumerate(bullets):
        p = body.paragraphs[0] if index == 0 else body.add_paragraph()
        p.text = text; p.font.name = "Microsoft YaHei"; p.font.size = PptPt(20)
        p.space_after = PptPt(14); p.level = 0


def make_presentation(target: Path) -> None:
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    add_slide(prs, "监盘通", ["基于 MCP 的资产现场盘点与数字取证平台", "数字科创大赛参赛作品"])
    add_slide(prs, "业务痛点", ["网页、底稿、表格与手机相机反复切换", "弱网下容易产生不完整清单", "同场多项设备逐项拍照，重复劳动明显", "照片、资产与 PDF 依赖人工整理"])
    add_slide(prs, "产品流程", ["MCP 安全读取项目与盘点资产", "完整快照落地，现场可离线工作", "统一清单筛选、搜索与批量选择", "单项或多项共用照片盘点", "自动生成 PDF 与项目 ZIP"])
    add_slide(prs, "核心创新", ["标准 MCP 替代 WebView 页面抓取", "完整快照与稳定身份保护历史证据", "一次拍摄关联多项资产，分别形成成果", "线上权威数据与本地现场修改分离"])
    add_slide(prs, "安全与可靠性", ["只读 MCP 工具白名单", "Token 经 Android Keystore 与 AES-GCM 加密", "失败不覆盖旧凭据和完整清单", "当前版本不回写业务系统", "自动测试、构建、Release 与 SHA-256"])
    add_slide(prs, "现场演示", ["连接竞赛专用 Token", "同步测试项目 20250715002", "筛选待盘点资产并查看元数据", "批量选择同场设备，共用一次拍摄", "查看各资产 PDF，断网复核并导出 ZIP"])
    add_slide(prs, "应用价值", ["减少数据搬运与重复拍摄", "提升弱网环境下的连续作业能力", "保持证据与资产关系可追溯", "为结构化成果上传和智能复核提供基础"])
    prs.save(target)


if __name__ == "__main__":
    markdown_to_docx(DOCS / "作品说明书.md", DOCS / "监盘通-作品说明书.docx")
    markdown_to_docx(DOCS / "可复现材料.md", DOCS / "监盘通-可复现材料.docx")
    make_presentation(DOCS / "监盘通-演示材料.pptx")
